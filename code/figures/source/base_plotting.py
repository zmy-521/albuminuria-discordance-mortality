# Release adaptation: paths only unless noted in docs/RELEASE_CHANGES.md.
import os as _release_os
def _release_path(env, relative=''):
    root = _release_os.environ.get(env)
    if not root: raise RuntimeError('Set environment variable '+env+'; see docs/REPRODUCIBILITY.md')
    return _release_os.path.join(root, relative)
from pathlib import Path
import json, csv, hashlib, shutil
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.lines import Line2D
from matplotlib.ticker import FixedLocator, FixedFormatter, NullLocator
from openpyxl import load_workbook

ROOT = Path(_release_path('ODKD_FIGURE_ROOT'))
MAIN = ROOT/'Main_candidates'
SUPP = ROOT/'Supplement_candidates'
for p in (MAIN,SUPP,ROOT/'Data'): p.mkdir(exist_ok=True)
MAINFILE = ROOT/'Inputs/19_final_clinical_results.xlsx'
EXTRAFILE = ROOT/'Inputs/02_supplement_statistical_results.xlsx'
BOOKS = {'main':load_workbook(MAINFILE,data_only=True),'extra':load_workbook(EXTRAFILE,data_only=True)}
USED=[]

def rows(book,sheet):
    s=BOOKS[book][sheet]
    heads=[c.value for c in s[4]]
    out=[]
    for i, vals in enumerate(s.iter_rows(min_row=5,values_only=True),5):
        if vals[0] is None: continue
        d={h:v for h,v in zip(heads,vals) if h is not None}
        d['_workbook']=MAINFILE.name if book=='main' else EXTRAFILE.name
        d['_sheet']=sheet; d['_row']=i
        out.append(d)
    return out

def one(book,sheet,**conditions):
    rr=[r for r in rows(book,sheet) if all(r.get(k)==v for k,v in conditions.items())]
    assert len(rr)==1,(sheet,conditions,len(rr))
    USED.append(rr[0]);return rr[0]

plt.rcParams.update({'font.family':'sans-serif','font.sans-serif':['DejaVu Sans'],
 'font.size':11,'axes.titlesize':12,'axes.labelsize':11,'xtick.labelsize':10,
 'ytick.labelsize':11,'text.color':'#20282E','axes.labelcolor':'#20282E',
 'axes.edgecolor':'#7D858B','axes.linewidth':.8,'savefig.facecolor':'white',
 'figure.facecolor':'white','svg.fonttype':'none','axes.spines.top':False,
 'axes.spines.right':False,'pdf.fonttype':42})
C={'G0':'#687581','G1':'#B38534','G2':'#B3473A','G3':'#287C8E'}
LIGHT={'G0':'#F0F3F5','G1':'#F8F1E5','G2':'#FBEDEA','G3':'#EAF4F6'}
GROUPS=['G0','G1','G2','G3']
NAMES={'G0':'Both below threshold','G1':'Concentration only','G2':'UACR only','G3':'Both above threshold'}
INK='#20282E'; MUTED='#5F6B72'; GRID='#E6EAED'

def header(fig,title,subtitle):
    fig.text(.055,.958,title,fontsize=18,weight='bold',ha='left',va='top')
    fig.text(.055,.908,subtitle,fontsize=10.5,color=MUTED,ha='left',va='top')

def panel(fig,x,y,letter,title):
    fig.text(x,y,letter,fontsize=14,weight='bold',va='top')
    fig.text(x+.029,y-.001,title,fontsize=12,weight='bold',va='top')

def note(fig,text,y=.048):
    fig.text(.055,y,text,fontsize=9,color=MUTED,va='bottom',linespacing=1.5)

def tidy(ax,grid='x'):
    ax.set_axisbelow(True)
    ax.grid(axis=grid,color=GRID,lw=.65)
    ax.spines['left'].set_visible(False)
    ax.tick_params(axis='y',length=0)

def logx(ax,ticks,limits):
    ax.set_xscale('log'); ax.set_xlim(*limits)
    ax.xaxis.set_major_locator(FixedLocator(ticks))
    ax.xaxis.set_major_formatter(FixedFormatter([f'{t:g}' for t in ticks]))
    ax.xaxis.set_minor_locator(NullLocator())
    ax.axvline(1,color='#8B9298',lw=1,ls=(0,(3,3)),zorder=1)
    tidy(ax)

def interval(ax,y,r,color=INK,marker='o',openmark=False,keys=('HR','low','high'),scale=1):
    e,l,h=[r[k]*scale for k in keys]
    assert np.isfinite([e,l,h]).all() and l<=e<=h,(r,e,l,h)
    ax.errorbar(e,y,xerr=[[e-l],[h-e]],fmt=marker,ms=7.3,mew=1.5,
      mfc='white' if openmark else color,mec=color,ecolor=color,elinewidth=1.55,capsize=3.4,zorder=3)

def fmt(r,dec=2,keys=('HR','low','high'),scale=1):
    return f"{r[keys[0]]*scale:.{dec}f} ({r[keys[1]]*scale:.{dec}f}–{r[keys[2]]*scale:.{dec}f})"

def save(fig,name,supp=False):
    fig.savefig((SUPP if supp else MAIN)/(name+'.png'),dpi=240)
    plt.close(fig)

def classification():
    fig=plt.figure(figsize=(12,8.0))
    header(fig,'Classification discordance in two data sources',
           'Thresholds: urinary albumin concentration 30 mg/L; UACR 30 mg/g')
    panel(fig,.055,.848,'A','NHANES 1999–2018 · n = 5,905')
    panel(fig,.54,.848,'B','Wanbei clinical sample · n = 2,261')
    fig.text(.055,.803,'Adults with diabetes; eGFRcr ≥60 mL/min/1.73 m²',fontsize=9.5,color=MUTED)
    for x,booksource,total in [(.125,'NHANES',5905),(.61,'Wanbei',2261)]:
        ax=fig.add_axes([x,.49,.31,.275])
        ax.set_xlim(0,2);ax.set_ylim(0,2);ax.invert_yaxis()
        ax.set_xticks([.5,1.5],['UACR <30','UACR ≥30'])
        ax.xaxis.tick_top();ax.tick_params(axis='x',length=0,pad=9)
        ax.set_yticks([.5,1.5],['Alb <30','Alb ≥30']);ax.tick_params(axis='y',length=0,pad=9)
        for sp in ax.spines.values():sp.set_visible(False)
        for g,col,row in [('G0',0,0),('G2',1,0),('G1',0,1),('G3',1,1)]:
            r=one('main','Group counts',threshold_mgl=30,group=g) if booksource=='NHANES' else one('main','Hospital groups',source='Wanbei',threshold=30,group=g)
            n=r['n']
            ax.add_patch(Rectangle((col+.012,row+.012),.976,.976,facecolor=LIGHT[g],edgecolor=C[g] if g=='G2' else '#D4DBDF',lw=1.8 if g=='G2' else .8))
            ax.text(col+.5,row+.23,g,color=C[g],ha='center',va='center',weight='bold',fontsize=12)
            ax.text(col+.5,row+.51,f'{n:,}',ha='center',va='center',weight='bold',fontsize=20)
            ax.text(col+.5,row+.78,f'{n/total*100:.2f}% of sample',ha='center',va='center',fontsize=10,color=MUTED)
    fig.text(.055,.459,'A–B: observed counts and unweighted percentages; cell area is not proportional to sample size.',fontsize=9,color=MUTED)
    panel(fig,.055,.435,'C','G2 frequency under two different denominators')
    ax=fig.add_axes([.34,.145,.42,.22])
    ax.set_xlim(0,30);ax.set_ylim(-.5,2.8)
    ax.set_yticks([2,0],['Among concentration-negative\nparticipants (G0 + G2)','Among UACR-positive\nparticipants (G2 + G3)'])
    ax.tick_params(axis='y',pad=15);ax.set_xlabel('G2 proportion (%)');tidy(ax)
    for y,den in [(2,'concentration_negative'),(0,'ACR_positive')]:
        r=one('main','Classification',source='NHANES5905',threshold=30,denominator=den)
        h=one('main','Hospital classification',source='Wanbei',threshold=30,denominator=den)
        interval(ax,y+.22,r,C['G2'],keys=('weighted','low','high'),scale=100)
        ax.plot(h['fraction']*100,y-.22,'s',ms=7,mfc='white',mec=INK,mew=1.3)
        ax.text(1.055,y+.22,f"{r['weighted']*100:.2f}% ({r['low']*100:.2f}–{r['high']*100:.2f})",transform=ax.get_yaxis_transform(),va='center',fontsize=10,color=C['G2'])
        ax.text(1.055,y-.22,f"{h['fraction']*100:.2f}%  [{h['numerator']}/{h['n']:,}]",transform=ax.get_yaxis_transform(),va='center',fontsize=10)
    ax.legend(handles=[Line2D([0],[0],color=C['G2'],marker='o',lw=1.5,label='NHANES: survey-weighted, 95% CI'),Line2D([0],[0],color=INK,marker='s',mfc='white',lw=0,label='Wanbei: unweighted')],loc='lower left',bbox_to_anchor=(-.63,1.01),frameon=False,ncol=2,fontsize=9.5)
    note(fig,'Alb = urinary albumin concentration. The concentration cutoff is an operational research threshold.\nWanbei data describe classification only; no external mortality validation. No hospital confidence intervals were supplied.',y=.026)
    save(fig,'Candidate_01_Classification')

def main_mortality():
    fig=plt.figure(figsize=(12,6.8))
    header(fig,'Directional classification and long-term mortality',
           'NHANES 1999–2018 · Survey-weighted Cox Model B · n = 5,155; deaths = 1,115')
    ys={'G0':4.8,'G1':3.8,'G2':2.8,'G3':1.8}; direct=.25
    ax=fig.add_axes([.58,.22,.235,.58]);ax.set_ylim(-.35,5.5)
    logx(ax,[.5,1,2,3],(.5,3.25));ax.set_yticks([]);ax.set_xlabel('Hazard ratio (log scale)')
    trans=ax.get_yaxis_transform()
    for x,t in [(-2.23,'Classification'),(-.68,'n'),(-.30,'Deaths'),(1.12,'HR (95% CI)')]:ax.text(x,5.42,t,transform=trans,weight='bold',fontsize=10)
    for g in GROUPS:
        y=ys[g];nr=one('main','Group counts',threshold_mgl=30,group=g)
        ax.text(-2.23,y,f'{g}  {NAMES[g]}',transform=trans,va='center',color=C[g],weight='bold' if g=='G2' else 'normal')
        ax.text(-.68,y,f"{nr['ccB']:,}",transform=trans,va='center')
        ax.text(-.30,y,f"{nr['death_ccB']:,}",transform=trans,va='center')
        if g=='G0':
            ax.plot(1,y,'D',color=C[g],ms=6);value='1.00 (reference)'
        else:
            r=one('main','Cox evidence',scenario='primary',model='B',design='survey_Taylor',contrast=f'{g}_G0')
            interval(ax,y,r,C[g]);value=fmt(r)
        ax.text(1.12,y,value,transform=trans,va='center',fontsize=11)
    ax.axhline(1.02,xmin=0,xmax=1,color=GRID,lw=1)
    ax.text(-2.23,.72,'Direct comparison of discordant groups',transform=trans,weight='bold',fontsize=10,color=MUTED)
    r=one('main','Cox evidence',scenario='primary',model='B',design='survey_Taylor',contrast='G2_G1')
    interval(ax,direct,r,INK,'D',True)
    ax.text(-2.23,direct,'G2 vs G1 (G1 reference)',transform=trans,va='center')
    ax.text(1.12,direct,fmt(r),transform=trans,va='center',fontsize=11)
    note(fig,'Model B: age spline, sex, race/ethnicity, survey cycle, BMI, HbA1c, eGFRcr, smoking, hypertension, prior CVD and PIR.\nG2 vs G0 is the original primary comparison. Counts above match the complete-case Cox population.\nG0: Alb−/UACR−; G1: Alb+/UACR−; G2: Alb−/UACR+; G3: Alb+/UACR+. Thresholds: 30 mg/L and 30 mg/g.',y=.04)
    save(fig,'Candidate_02_LongTerm_Mortality')

def components():
    fig=plt.figure(figsize=(12,8.4))
    header(fig,'Urine components and the within-A2 comparison',
           'Post hoc explanatory analyses · Frozen Model B covariates · Survey-weighted Cox estimates')
    panel(fig,.055,.85,'A','Conditional component associations · n = 5,155; deaths = 1,115')
    ax=fig.add_axes([.50,.55,.26,.24]);ax.set_ylim(-.55,2.55)
    logx(ax,[.65,.8,1,1.25],(.62,1.29));ax.set_yticks([]);ax.set_xlabel('HR per doubling (log scale)')
    labels=[('Alb_per_doubling','Albumin doubles; creatinine fixed',2,C['G3']),('UCr_per_doubling_at_fixed_Alb','Creatinine doubles; albumin fixed',1,C['G0']),('UCr_per_doubling_at_fixed_ACR','Creatinine and albumin both double;\nUACR fixed',0,C['G2'])]
    for key,label,y,col in labels:
        r=one('extra','Effects',scenario='original_linear_eGFR',contrast=key)
        interval(ax,y,r,col)
        ax.text(-1.70,y,label,transform=ax.get_yaxis_transform(),va='center',fontsize=11,color=col)
        ax.text(1.12,y,fmt(r,3),transform=ax.get_yaxis_transform(),va='center',fontsize=11)
    ax.text(1.12,2.55,'HR (95% CI)',transform=ax.get_yaxis_transform(),weight='bold',fontsize=10)
    fig.text(.055,.485,'Fixed-UACR contrast: H₀: βAlb + βCr = 0; one-degree-of-freedom Wald P = 0.000253.',fontsize=10,color=MUTED)
    panel(fig,.055,.423,'B','Within A2: G2 vs G3-A2')
    fig.text(.084,.382,'UACR 30–<300 mg/g; adjusted for log₂ UACR\nn = 1,120; deaths = 350',fontsize=10,color=MUTED,linespacing=1.45,va='top')
    ax=fig.add_axes([.12,.218,.31,.095]);ax.set_ylim(-.7,.7);ax.set_yticks([])
    logx(ax,[.7,1,1.5,2],(.65,2.05));ax.set_xlabel('Hazard ratio (log scale)')
    r=one('extra','Effects',scenario='original_linear_eGFR',contrast='G2_vs_G3_A2_adjust_log2ACR')
    interval(ax,0,r,C['G2']);ax.text(.5,1.05,f'{fmt(r,3)}; P = {r["p"]:.3f}',transform=ax.transAxes,ha='center',va='bottom',fontsize=11,weight='bold')
    fig.text(.084,.135,'G2: 222 participants / 76 deaths\nG3-A2: 898 participants / 274 deaths',fontsize=10,color=MUTED,linespacing=1.5,va='top')
    panel(fig,.55,.423,'C','Observed UACR support in the A2 model')
    ax=fig.add_axes([.64,.155,.285,.205]);bins=['30-<60','60-<100','100-<200','200-<300']
    ax.set_xlim(0,92);ax.set_ylim(-.55,3.55);ax.invert_yaxis()
    for g,n,off in [('G2',222,-.16),('G3',898,.16)]:
        for y,b in enumerate(bins):
            r=one('extra','A2 support bins',population='A2_B',group=g,interval=b)
            p=r['n']/n*100
            ax.barh(y+off,p,height=.26,color=C[g],alpha=.87)
            ax.text(p+1.5,y+off,f"{p:.1f}% ({r['n']})",va='center',fontsize=9,color=C[g])
    ax.set_yticks(range(4),[b.replace('-<','–<') for b in bins]);ax.set_xlabel('Within-group percentage (unweighted)')
    ax.set_ylabel('UACR (mg/g)',fontsize=10);tidy(ax)
    ax.legend(handles=[Rectangle((0,0),1,1,fc=C['G2'],label='G2'),Rectangle((0,0),1,1,fc=C['G3'],label='G3-A2')],loc='lower left',bbox_to_anchor=(0,1.01),ncol=2,frameon=False,fontsize=9)
    note(fig,'A and B estimate different contrasts; the A2 result is not an independent validation of the component model.\nSame UACR implies proportional changes in both components. Associations are model-dependent and not causal.',y=.032)
    save(fig,'Candidate_03_Components_A2')

def absolute_risk():
    fig=plt.figure(figsize=(12,7.3))
    header(fig,'Five-year absolute mortality risk',
           'NHANES baseline 1999–2014 · Standardized to 3,789 complete-case participants · Both frozen methods retained')
    panel(fig,.055,.84,'A','Standardized risk by classification')
    panel(fig,.535,.84,'B','Between-group risk differences')
    ax=fig.add_axes([.095,.24,.345,.47]);ax.set_ylim(-.65,3.65);ax.invert_yaxis();ax.set_xlim(0,17.3)
    ax.set_yticks(range(4),GROUPS);ax.set_xlabel('Five-year mortality risk (%)');tidy(ax)
    ax2=fig.add_axes([.66,.24,.245,.47]);ax2.set_ylim(-.55,2.9);ax2.invert_yaxis();ax2.set_xlim(-6.3,9)
    ax2.set_yticks([.5,2],['G2 − G0','G2 − G1']);ax2.tick_params(axis='y',pad=14);ax2.set_xlabel('Risk difference (percentage points)');tidy(ax2)
    ax2.axvline(0,color='#8B9298',ls=(0,(3,3)),lw=1)
    for method,offset,marker,col,opened in [('logistic',-.17,'o',INK,False),('cox',.17,'s','#637785',True)]:
        for y,g in enumerate(GROUPS):
            key=f'30_known_risk_{g}' if method=='logistic' else f'AdjustedB_{g}'
            r=one('main','Absolute risk',estimand=key)
            interval(ax,y+offset,r,col,marker,opened,('estimate','low','high'),100)
            ax.text(r['high']*100+.32,y+offset,f"{r['estimate']*100:.2f}",va='center',fontsize=9.5,color=col)
        for y,contrast in [(.5,'G2_G0'),(2,'G2_G1')]:
            key=f'30_known_RD_{contrast}' if method=='logistic' else f'AdjustedB_RD_{contrast}.AdjustedB_G2'
            r=one('main','Absolute risk',estimand=key)
            interval(ax2,y+offset,r,col,marker,opened,('estimate','low','high'),100)
            ax2.text(.02,y+offset+.18,fmt(r,2,('estimate','low','high'),100),transform=ax2.get_yaxis_transform(),fontsize=9,va='top',color=col)
    handles=[Line2D([0],[0],color=INK,marker='o',lw=1.5,label='Fixed-horizon logistic (primary)'),Line2D([0],[0],color='#637785',marker='s',mfc='white',lw=1.5,label='Cox standardization (sensitivity)')]
    fig.legend(handles=handles,loc='upper left',bbox_to_anchor=(.066,.793),frameon=False,ncol=2,fontsize=10)
    note(fig,'Points and bars show estimates and 95% CIs. Logistic fitting: 3,785 participants with known 60-month outcome;\nstandardization target: 3,789. Cox uses the frozen 1999–2014 complete-case analysis.\nRisk-difference CIs are read directly from the joint estimates, not obtained by subtracting separate risk CIs.',y=.056)
    save(fig,'Candidate_04_FiveYear_Risk')

def flow():
    fig=plt.figure(figsize=(12,7.7))
    header(fig,'Urine flow and short-interval excretion',
           'NHANES 2009–2018 · Evaluable subset n = 3,261 / 3,539 · Survey-weighted arithmetic means and 95% CIs')
    settings=[('flow','Urine flow','mL/min',(0,2.4),'A'),('CER','Estimated creatinine excretion','mg/min',(0,1.38),'B'),('AER','Estimated albumin excretion','µg/min · log scale',(5,550),'C')]
    for i,(variable,title,unit,lims,letter) in enumerate(settings):
        x=.055+i*.318
        panel(fig,x,.835,letter,title)
        ax=fig.add_axes([x+.036,.285,.244,.46]);ax.set_xlim(-.5,3.5);ax.set_ylim(*lims)
        if variable=='AER':
            ax.set_yscale('log');ax.yaxis.set_major_locator(FixedLocator([5,10,30,100,300]));ax.yaxis.set_major_formatter(FixedFormatter(['5','10','30','100','300']));ax.yaxis.set_minor_locator(NullLocator())
        for j,g in enumerate(GROUPS):
            r=one('extra','Flow characteristics',population=g,variable=variable)
            e,l,h=r['mean'],r['low'],r['high'];assert l<=e<=h
            ax.errorbar(j,e,yerr=[[e-l],[h-e]],fmt='o',ms=8,color=C[g],capsize=4,elinewidth=1.6)
            label=f'{e:.2f}' if variable!='AER' else f'{e:.1f}'
            ax.annotate(label,(j,h),xytext=(0,9),textcoords='offset points',ha='center',fontsize=11,color=C[g])
        ax.set_xticks(range(4),GROUPS);ax.set_ylabel(unit);tidy(ax,'y');ax.spines['left'].set_visible(True)
        for tick,g in zip(ax.get_xticklabels(),GROUPS):tick.set_color(C[g]);tick.set_weight('bold' if g=='G2' else 'normal')
    fig.text(.091,.235,'G0: n = 2,276       G1: n = 176       G2: n = 123       G3: n = 686',fontsize=11)
    r=one('extra','G2 AER20',group='G2')
    fig.text(.091,.185,f"G2 with estimated AER ≥20 µg/min: {r['weighted']*100:.1f}% (95% CI {r['low']*100:.1f}–{r['high']*100:.1f}); observed count 59/123.",fontsize=10.5,color=C['G2'])
    note(fig,'AER and CER use the same urine analytes and collection flow; they are not independent diagnostic reference measures.\nShort-interval estimates are not 24-hour excretion rates. Missing flow was not imputed (278 participants).\nG0: both below threshold; G1: concentration only; G2: UACR only; G3: both above threshold.',y=.047)
    save(fig,'Candidate_05_Urine_Flow')

def sensitivity():
    fig=plt.figure(figsize=(12,9.2))
    header(fig,'Sensitivity analyses of the G2 versus G0 association',
           'Supplement candidate · Frozen survey-weighted estimates · Each row retains its own analytic population')
    records=[]
    def add(label,r):records.append((label,r))
    add('Original Model B',one('main','Cox evidence',scenario='primary',model='B',design='survey_Taylor',contrast='G2_G0'))
    add('Albumin cutoff 20 mg/L',one('main','Cox evidence',scenario='threshold_20',model='B',design='survey_Taylor',contrast='G2_G0'))
    add('Negative domain + log₂ albumin',one('main','Stage3 Cox',analysis='negative_domain_B_log2alb',contrast='G2_G0'))
    add('24-month landmark',one('main','Stage3 Cox',analysis='landmark24',contrast='G2_G0'))
    add('eGFRcr spline (3 df)',one('extra','Effects',scenario='fixed_RCS3_eGFR',contrast='original_G2_G0_RCS_check'))
    add('PIR omitted; expanded sample',one('main','Cox evidence',scenario='threshold_30',model='NoPIR_expanded',design='survey_Taylor',contrast='G2_G0'))
    add('Complete-case selection weighting',one('main','Selection sensitivity',contrast='G2_G0'))
    add('Baseline cycles 2007–2018',one('main','Cox evidence',scenario='post2007',model='B',design='survey_Taylor',contrast='G2_G0'))
    add('Unharmonized urine measurements',one('main','Cox evidence',scenario='urine_unadjusted',model='B',design='survey_Taylor',contrast='G2_G0'))
    ax=fig.add_axes([.56,.21,.245,.61]);ax.set_ylim(-.65,len(records)-.1);ax.invert_yaxis()
    logx(ax,[.75,1,1.5,2,3,4],(.7,4.05));ax.set_yticks([]);ax.set_xlabel('Hazard ratio (log scale)')
    for x,t in [(-2.04,'Analysis'),(-.46,'n / deaths'),(1.12,'HR (95% CI)')]:ax.text(x,-.5,t,transform=ax.get_yaxis_transform(),weight='bold',fontsize=10)
    for i,(label,r) in enumerate(records):
        interval(ax,i,r,C['G2'] if i==0 else INK)
        ax.text(-2.04,i,label,transform=ax.get_yaxis_transform(),va='center',fontsize=10.5,weight='bold' if i==0 else 'normal')
        ax.text(-.46,i,f"{r['n']:,} / {r['events']:,}",transform=ax.get_yaxis_transform(),va='center',fontsize=9.5)
        ax.text(1.12,i,fmt(r),transform=ax.get_yaxis_transform(),va='center',fontsize=10)
    note(fig,'These analyses are not independent replications. The 20 mg/L definition changes group membership; landmark analysis\nis conditional on survival and observation to 24 months. Selection weighting cannot remove all missingness bias.\nThe urine-measurement comparison uses the frozen unharmonized values and does not replace the primary harmonized analysis.',y=.063)
    save(fig,'Supplement_01_Sensitivity',True)

if __name__=='__main__':
    classification();main_mortality();components();absolute_risk();flow();sensitivity()
    unique={(r['_workbook'],r['_sheet'],r['_row']):r for r in USED}
    provenance={'input_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [MAINFILE,EXTRAFILE]},'source_rows':list(unique.values()),'status':'Review candidates; not user-approved; no models fitted.'}
    (ROOT/'Data/figure_source_rows.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2))
    with (ROOT/'Data/figure_source_cells.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['Workbook','Sheet','Excel row','Field','Frozen value'])
        for r in unique.values():
            for k,v in r.items():
                if not k.startswith('_'):w.writerow([r['_workbook'],r['_sheet'],r['_row'],k,v])
    print(json.dumps({'figures':6,'unique_source_rows':len(unique),'output':str(ROOT)},ensure_ascii=False))
