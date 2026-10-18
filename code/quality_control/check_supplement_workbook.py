from pathlib import Path
import json,math
from openpyxl import load_workbook
from PIL import Image,ImageDraw
S=Path(__import__('os').environ['ODKD_WORK_ROOT'])/'Review_supplement_urine_creatinine';p=json.loads((S/'internal/workbook_payload.json').read_text(encoding='utf8'));w=load_workbook(S/'02_supplement_statistical_results.xlsx',data_only=True);n=0
for t in p:
 sh=w[t['name']]
 for i,row in enumerate(t['rows'],5):
  for j,v in enumerate(row,1):
   a=sh.cell(i,j).value
   assert a==v or isinstance(v,(float,int)) and isinstance(a,(float,int)) and math.isclose(a,v,rel_tol=1e-11,abs_tol=1e-12),(t['name'],i,j,a,v)
   if t['headers'][j-1].lower() in ['p','f_p','p_value','pvalue']:assert a!=0
   n+=1
print('PASS numerical cells',n)
